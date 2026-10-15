"""Gera o relatório em Word e PDF a partir de resultados/resumo.json.

    python relatorio/gerar_relatorio.py

Regra de redação do usuário: português, sem datas, sem travessão e sem dois-pontos
na prosa. A função `checar` recusa qualquer parágrafo que viole isso.
"""
import json
import subprocess
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

RAIZ = Path(__file__).resolve().parents[1]
RES = RAIZ / "resultados"
SAIDA = RAIZ / "relatorio"


def checar(t: str):
    proibidos = [":", "—", "–"]
    for c in proibidos:
        if c in t:
            raise ValueError(f"caractere proibido {c!r} no texto: {t[:120]}")
    return t


def br(x, casas=1):
    """Número com vírgula decimal."""
    if x is None:
        return "n/d"
    s = f"{x:.{casas}f}"
    return s.replace(".", ",")


def sci(x, casas=1):
    if x is None:
        return "n/d"
    m, e = f"{x:.{casas}e}".split("e")
    return f"{m.replace('.', ',')}e{int(e)}"


class Rel:
    def __init__(self):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10.5)
        for s in self.d.sections:
            s.left_margin = s.right_margin = Cm(2.2)
        self.nfig = 0
        self.ntab = 0

    def titulo(self, t, nivel=1):
        self.d.add_heading(checar(t), level=nivel)

    def p(self, t, negrito=False):
        par = self.d.add_paragraph()
        r = par.add_run(checar(t))
        r.bold = negrito
        par.paragraph_format.space_after = Pt(4)
        return par

    def lista(self, itens):
        for i in itens:
            self.d.add_paragraph(checar(i), style="List Bullet")

    def figura(self, arquivo, legenda, largura=15.5):
        if not Path(arquivo).exists():
            self.p(f"Figura ausente ({Path(arquivo).name}). Rode exemplos/gerar_resultados.py.")
            return
        self.nfig += 1
        self.d.add_picture(str(arquivo), width=Cm(largura))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        par = self.d.add_paragraph()
        r = par.add_run(f"Figura {self.nfig}. ")
        r.bold = True
        par.add_run(checar(legenda))
        par.runs[-1].font.size = Pt(9)

    def tabela(self, cab, linhas, legenda):
        self.ntab += 1
        par = self.d.add_paragraph()
        r = par.add_run(f"Tabela {self.ntab}. ")
        r.bold = True
        par.add_run(checar(legenda))
        t = self.d.add_table(rows=1, cols=len(cab))
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, cab):
            c.text = checar(str(h))
        for ln in linhas:
            cells = t.add_row().cells
            for c, v in zip(cells, ln):
                c.text = checar(str(v))
        for row in t.rows:
            for c in row.cells:
                for pp in c.paragraphs:
                    for rr in pp.runs:
                        rr.font.size = Pt(8.5)
        self.d.add_paragraph()

    def salvar(self, caminho):
        self.d.save(caminho)


def intervalo(m, chave, f=lambda v: v, casas=1, cientifico=False):
    v = m[chave]
    fmt = sci if cientifico else br
    return f"{fmt(f(v['mediana']), casas)} ({fmt(f(v['p2.5']), casas)} a {fmt(f(v['p97.5']), casas)})"


def main():
    R = json.loads((RES / "resumo.json").read_text())
    rel = Rel()
    rel.d.add_heading("Simulação da porta local implantada do GFET", 0)
    rel.p("Relatório técnico do pacote gfetsim. Laboratório de Microeletrônica, PSI, Escola Politécnica da USP.")

    # ------------------------------------------------------------------ resumo
    im, rc, ox, el = R["implantacao"], R["recozimento"], R["oxidacao"], R["eletrico"]
    t20A = {d["T_C"]: d["t_min"] for d in ox["t20"]["A|ho_plummer|acumulo_classico"]}
    t20s = {d["T_C"]: d["t_min"] for d in ox["t20"]["A|sem_aceleracao|acumulo_classico"]}
    mref = R["montecarlo"]["A|ho_plummer|acumulo|p"]
    vd_lo, vd_hi = abs(mref["V_D"]["p2.5"]), abs(mref["V_D"]["p97.5"])
    perda_op = mref["perda_si_5e12"]["mediana"]
    rel.titulo("Resumo")
    rel.p(
        "O gfetsim simula em cadeia a fabricação da porta local do GFET, da implantação de P a 30 keV até a "
        "capacitância e a curva de transferência do grafeno sobre 20 nm de SiO2. Cada número físico vem de um único "
        "arquivo de parâmetros, com fonte conferida e etiqueta de origem, e as discordâncias da literatura viram ramos "
        "separados que se propagam até o resultado. "
        f"A implantação dá Rp de {br(im['Rp_nm'])} nm e pico de {sci(im['pico_cm3'])} cm-3, com camada amorfa contínua "
        "a partir da superfície e a máscara de 150 nm bloqueando todo o feixe. "
        f"Após a variante A a porta n+ fica totalmente ativa, com resistência de folha de {br(rc['A (RTA)']['Rs_ohm_sq'])} ohm por quadrado. "
        f"O tempo para 20 nm de óxido seco a 900 C vai de {br(t20A[900.0])} min, com a aceleração de Ho e Plummer, a "
        f"{br(t20s[900.0])} min sem aceleração, e essa diferença é a maior incerteza do processo. "
        f"No lado elétrico, a perda de capacitância para o Si n+ fica perto de {br(100 * perda_op, 1)} % com 5e12 cm-2 "
        "no grafeno, bem menor que a estimada antes por depleção, e o risco principal passa a ser a posição do ponto de "
        f"Dirac, que fica entre {br(vd_lo)} e {br(vd_hi)} V em módulo (intervalo de 95 %) se a dopagem residual do grafeno "
        "for a que o Raman do artigo sugere.")

    # ------------------------------------------------------------------ escopo
    rel.titulo("O que o pacote faz e como ler os números")
    rel.p("O pacote tem seis módulos encadeados. Implantação (RustBCA e Pearson), recozimento (SPE, Fair e Tsai, "
          "solubilidade de Solmi, desativação), oxidação (Massoud, Ho e Plummer, segregação com fronteira móvel), "
          "elétrico (Poisson com Fermi-Dirac, grafeno com poças de carga, modelos de Kim e de Jiménez, ruptura e "
          "Fowler-Nordheim), incerteza (Monte Carlo por ramo e Sobol) e calibração (mínimos quadrados ponderados e emcee).")
    rel.p("Cada parâmetro tem etiqueta. [P] é publicado e conferido, [C] é cálculo próprio a partir de [P], [I] é "
          "inferência ou valor lembrado sem fonte conferida e [B] é dado de bancada. Valores [I] só entram quando a "
          "receita aceita os provisórios, e a lista completa dos que entraram está no fim deste relatório. Os testes de "
          "validação usam dados publicados com tolerância declarada antes de rodar. Quando um módulo não validou, isso "
          "está dito na seção dele.")
    rel.p("A rede da sessão não alcança as editoras. A conferência foi feita pelos resumos e, quando havia texto aberto, "
          "pelo texto completo (arXiv e teses da TU Wien). Os números que só existem em texto fechado ficaram [I].")

    # ------------------------------------------------------------------ receita
    rel.titulo("Receita simulada")
    rel.tabela(["etapa", "condição", "origem"], [
        ["implantação", "P+ 30 keV, 3e15 cm-2, 7 graus, temperatura ambiente", "B"],
        ["variante A da implantação", "máscara de SiO2 de 150 nm com janela", "B"],
        ["variante B da implantação", "lâmina inteira", "B"],
        ["recozimento A", "550 C 30 min seguido de RTA a 1050 C por 10 s", "B"],
        ["recozimento B", "550 C 30 min seguido de forno a 850 a 900 C por 30 min em N2", "B"],
        ["óxido de porta térmico", "O2 seco, alvo de 20 nm, consumo de 0,44 nm de Si por nm", "B"],
        ["óxido de porta alternativo", "sputtering RF reativo a 50 W, 5,1 min (artigo do usuário)", "B"],
        ["substrato", "Si(100) tipo n, resistividade varrida de 0,001 a 10 ohm cm", "B"],
        ["grafeno", "CVD transferido com PMMA (artigo do usuário)", "B"],
    ], "Receita e variantes.")

    # ------------------------------------------------------------------ módulo 1
    rel.titulo("Módulo 1, implantação")
    a2 = im["amorfo_faixa_A2"]
    rel.p(f"O RustBCA (alvo amorfo, 1e5 íons, potencial Kr-C) dá Rp de {br(im['Rp_nm'])} nm com erro estatístico de "
          f"{br(im['erro_Rp_nm'], 2)} nm, desvio de {br(im['dRp_nm'])} nm, assimetria de {br(im['assimetria'], 2)} e "
          f"curtose de {br(im['curtose'], 2)}. Como a curtose fica abaixo de 3, o tipo IV de Pearson não existe para esses "
          f"momentos e o código usa o tipo I, que reproduz os quatro momentos. O pico é {sci(im['pico_cm3'])} cm-3. "
          f"Nenhum de 1e5 íons atravessou 150 nm de SiO2, então a transmissão da máscara da variante A é menor que 1e-5.")
    rel.p("O critério de amorfização por densidade de energia de Dennis e Hale (A4), 6e23 eV/cm3, prevê limiar de "
          "cerca de 1,5e14 cm-2 para P de 10 a 60 keV, dez vezes abaixo do limiar medido a 300 K, de 1 a 3e15 cm-2 (A2). "
          "Narayan et al. (A6) mostram que 6e23 eV/cm3 é o valor a 4 K sem recozimento dinâmico, e o próprio resumo de A4 "
          "diz que a temperatura mais alta exige incluir a difusão do dano. O modelo passou a usar o limiar medido a 300 K "
          f"(ramo limiar_TA). Com a faixa de A2, a camada amorfa vai da superfície até {br(a2[0][1] if a2[0] else 0)} nm "
          f"(limiar de 1e15) ou fica restrita ao pico (limiar de 3e15, igual à dose da receita).")
    rel.p("O alcance não tem validação forte. Grant, Williams e Dodds (G1) medem por RBS alcances 15 a 70 % acima da teoria "
          "LSS, e Kobayashi e Gibson (K1) medem P por NRA concordando com o TRIM. O Rp do RustBCA fica 15 a 21 % abaixo da "
          "relação de G1 na faixa em que ela vale, o que está dentro da discordância publicada. O Monte Carlo carrega por "
          "isso um ramo com Rp 18 % maior.")
    rel.figura(RES / "fig_implantacao.png", "Perfil de P implantado (pontos do RustBCA e ajuste de Pearson I) com a camada "
               "amorfa para o limiar central de A2, e densidade de energia nuclear depositada na dose da receita, com o "
               "critério de baixa temperatura de A4 em vermelho.")
    rel.tabela(["teste", "dado publicado", "resultado"], [
        ["Pearson recupera os momentos", "verificação numérica", "passou"],
        ["alcance dentro da discordância", "G1 e K1", "passou (teste fraco)"],
        ["critério de 6e23 eV/cm3 a 300 K", "A2, A4, A6", "falhou, documentado"],
        ["amorfo a 1e16 cm-2", "A3", "passou"],
        ["máscara e conservação da dose", "verificação", "passou"],
    ], "Validação do módulo 1.")

    # ------------------------------------------------------------------ módulo 2
    rel.titulo("Módulo 2, recozimento")
    rel.p("A camada amorfa recresce por SPE na etapa de 550 C em poucos minutos com o prefator provisório, de acordo com "
          "R14, que vê ativação máxima já a 550 C. O P incorporado na SPE fica ativo de forma metaestável. Na etapa "
          "seguinte, a difusão segue Fair e Tsai com campo elétrico, o P acima de C_sat(T) precipita e o ativo cai para "
          "n_e(T), ambos de Solmi et al. (N24), lidos no resumo.")
    linhas = []
    for k in ("A (RTA)", "A metaestável", "B 850 C", "B 900 C"):
        v = rc[k]
        linhas.append([k, br(v["fracao_ativa"], 2), br(v["Rs_ohm_sq"]), br(v["Rs_limite_Li_Han"]), sci(v["ativo_superficie_cm3"]),
                       br(v["xj_1e17_nm"], 0)])
    rel.tabela(["variante", "fração ativa", "Rs (ohm/q)", "Rs com 67 % ativo (R19)", "ativo na superfície (cm-3)",
                "profundidade em 1e17 (nm)"], linhas, "Resultado do recozimento antes da oxidação de porta.")
    rel.p("O modelo valida o recozimento em duas etapas. Para P a 50 keV com 550 C seguido de 1050 C por 10 s, prevê "
          "ativação dentro da faixa de 75 a 100 % de Yu et al. (R13). O modelo não valida o recozimento em uma etapa. "
          "Prevê 100 % onde R13 mede 40 a 60 % e onde Li e Han (R19) medem 67 % a 850 C. Falta a física dos aglomerados "
          "P-defeito (R12), cuja cinética só está no texto completo de R12b. Como a receita sempre passa pela etapa de "
          "550 C, a variante A fica coberta pela validação. A variante B não tem dado direto e por isso a tabela traz a "
          "resistência de folha que resultaria se só 67 % do P ficasse ativo, como limite pessimista.")
    rel.figura(RES / "fig_recozimento.png", "Perfis de P total após as variantes de recozimento. O ativo coincide com o "
               "total porque a difusão baixa o pico abaixo de n_e(T).", 11)

    # ------------------------------------------------------------------ módulo 3
    rel.titulo("Módulo 3, oxidação de porta")
    rel.p("Com 20 nm o óxido está no regime fino, abaixo dos 30 nm em que Deal e Grove foi validado (R25). O modelo usa a "
          "forma de Massoud com as constantes para Si(100) abaixo de 1000 C lidas em S1. A tabela de S1 escreve as "
          "constantes de tempo com expoente negativo, o que daria constantes de 1e-12 min. O código usa o expoente "
          "positivo, que dá minutos e decresce com a temperatura, como o efeito exige.")
    rel.p("A aceleração da oxidação pelo P é o ponto mais incerto. Ho e Plummer (R29, parâmetros em S2) preveem fator de "
          "3 a 5 no termo linear para 2 a 3e20 cm-3 na interface. Nomura et al. (R23) medem no máximo 2,1 em Si "
          "implantado. Massoud (N10) mostra que no início da oxidação seca de Si muito dopado o P acumula na interface "
          "e a taxa iguala a do Si leve, exatamente o regime de 20 nm. Os três ramos são mantidos separados.")
    T_lista = [800.0, 850.0, 900.0, 950.0, 1000.0]
    linhas = []
    for ra, nome in (("sem_aceleracao", "sem aceleração (N10)"), ("nomura", "F = 2,1 (R23)"), ("ho_plummer", "Ho e Plummer (R29)")):
        for var in ("A", "B900"):
            dd = {d["T_C"]: d["t_min"] for d in ox["t20"][f"{var}|{ra}|acumulo_classico"]}
            linhas.append([nome, var] + [br(dd[T], 1) for T in T_lista])
    linhas.append(["Deal e Grove sem aceleração", "fora da validade"] +
                  [br(ox["t20_cinetica_pura"]["deal_grove"][str(T)], 1) for T in T_lista])
    rel.tabela(["ramo de aceleração", "recozimento", "800 C", "850 C", "900 C", "950 C", "1000 C"], linhas,
               "Tempo em minutos para 20 nm de óxido seco sobre a porta n+, com acoplamento ao perfil de P.")
    rel.figura(RES / "fig_mapa_oxido.png", "Curvas de temperatura e tempo que dão 20 nm. As linhas contínuas e tracejadas "
               "usam fator de aceleração constante. Os pontos usam Ho e Plummer acoplado ao P ativo da interface.", 12)
    ri = ox["t20"]["A|ho_plummer|lado_oxido"][2]
    rel.p(f"A oxidação de 20 nm consome cerca de {br(ri['si_consumido_nm'])} nm de Si. No ramo em que o P acumula do lado "
          f"do óxido (R32, N12), cerca de {sci(ri['dose_no_oxido'])} cm-2 ficam no óxido, perto de 7 % da dose. No ramo "
          "em que o P acumulado fica inativo do lado do Si (N22, R20), o ativo nos primeiros 15 nm cai para 60 %.")
    rel.p("Este módulo não tem validação quantitativa. Nenhum resumo acessível traz espessura contra tempo de óxido seco "
          "fino em Si(100). Os testes conferem a solução numérica contra a analítica, o regime fino mais rápido (R26), a "
          "aceleração só no termo linear (R23), o ramo sem aceleração (N10), a conservação da dose e o consumo de Si. A "
          "recomendação é a mesma da seção 4.1 do artigo do usuário, medir o óxido sobre uma amostra de controle com a "
          "mesma implantação e o mesmo recozimento, e usar essa medida no módulo 6.")

    # ------------------------------------------------------------------ módulo 4
    rel.titulo("Módulo 4, elétrico")
    rel.p("A capacitância da pilha grafeno/SiO2/Si n+ vem de três termos em série, o óxido, a capacitância quântica do "
          "grafeno e a do Si. O Si entra pelo Poisson com Fermi-Dirac sobre o perfil ativo simulado. O grafeno entra com "
          "densidade de estados alargada por flutuação de potencial (R40), calibrada pelo n* de Ju e Lee (N17), o que dá "
          "capacitância quântica mínima finita no ponto de Dirac, como R39 mede.")
    pa = el["perdas_si"]
    linhas = []
    for nome in pa:
        linhas.append([nome] + [br(100 * pa[nome][k], 1) for k in ("2e+12", "5e+12", "1e+13", "-1e+13")])
    rel.tabela(["perfil do Si", "2e12 elétrons", "5e12 elétrons", "1e13 elétrons", "1e13 lacunas"], linhas,
               "Perda de capacitância para o Si, em por cento da capacitância com porta metálica, por densidade no grafeno (cm-2).")
    c11 = el["curvatura_1p1V"]
    txt = []
    for N in ("1e+19", "1e+20", "3e+20"):
        v = c11[N]
        txt.append(f"{N.replace('+', '')} cm-3 exige {sci(v['n_grafeno_cm2'])} cm-2 no grafeno e {br(v['E_ox_MVcm'])} MV/cm no óxido"
                   if v["E_ox_MVcm"] else f"{N.replace('+', '')} cm-3 não chega a 1,1 V na faixa calculada")
    eb = R["eletrico"]["ruptura_MVcm"]["adams_MVcm"]
    acima = [N.replace("+", "") for N in ("1e+19", "1e+20", "3e+20") if c11[N]["E_ox_MVcm"] and c11[N]["E_ox_MVcm"] > eb[0]]
    rel.p("O cálculo anterior do usuário (17 %, 6 % e 3,5 % para 1e19, 1e20 e 3e20 cm-3) é reproduzido exatamente pela "
          "aproximação de depleção com 1,1 V de curvatura de bandas. O Poisson completo mostra quanta carga essa curvatura "
          "exige. Para chegar a 1,1 V, " + ", ".join(txt) + ". " +
          (f"Com {' e '.join(acima)} cm-3 o campo passa da ruptura intrínseca de {br(eb[0], 0)} a {br(eb[1], 0)} MV/cm (B1), "
           "então a curvatura de 1,1 V nunca é alcançada e a perda por depleção superestima a real. " if acima else "") +
          "Nas densidades de operação, a perda real é a da Tabela, menos de 2 % para o perfil simulado.")
    rel.figura(RES / "fig_cv.png", "Capacitância normalizada em função da tensão de porta medida a partir do ponto de Dirac, "
               "para o perfil simulado e para dopagens uniformes. As linhas pontilhadas são as perdas do cálculo anterior por depleção.", 11)
    vd = el["V_D"]
    rel.p(f"O ponto de Dirac depende da dopagem residual do grafeno. O Raman do artigo (posição do G em 1592,6 cm-1 e "
          f"I(2D)/I(G) de 1,9) indica dopagem maior que 1,4e12 cm-2, comparando com Das et al. (D1), e menor que cerca de "
          f"1e13 cm-2. Sobre 20 nm de SiO2 isso coloca V_D entre {br(abs(vd['1.4e+12']))} e {br(abs(vd['1.0e+13']))} V em "
          "módulo, com sinal que o Raman não decide. A ruptura intrínseca de 10 a 13 MV/cm (B1, B2) corresponde a 20 a 26 V "
          "sobre 20 nm. Se a dopagem for da ordem de 1e13 cm-2, operar alguns volts além do ponto de Dirac já leva o óxido a "
          "metade do campo de ruptura. Medir V_D logo após a transferência, e recozer para dessorver água e resíduos antes "
          "de medir, é a forma de saber em qual extremo o dispositivo está.")
    fn = el["FN_A_cm2"]
    rel.p(f"A corrente de Fowler-Nordheim (massa efetiva de R45, barreira provisória de 3,1 eV) é de {sci(fn['6 MV/cm'])} "
          f"A/cm2 a 6 MV/cm e de {sci(fn['8 MV/cm'])} A/cm2 a 8 MV/cm, desprezível frente à corrente do canal abaixo de 6 MV/cm.")
    par = el["parasita_fora_do_n+"]
    rel.p("A resistividade do substrato não muda a C-V sob a porta, porque a camada n+ blinda o substrato. Ela importa "
          "fora da porta. Uma região n+ implantada num substrato tipo n não é isolada por junção, então todo o substrato "
          "fica no potencial da porta, e os contatos de fonte e dreno sobre 20 nm de óxido viram capacitores parasitas para "
          "a porta. Com substrato de 0,001 ohm cm a capacitância parasita é praticamente a do óxido. Com 10 ohm cm ela cai "
          f"para {br(100 * par['10.0']['C_min_sobre_Cox'], 0)} % do Cox na depleção, mas volta a cerca de "
          f"{br(100 * par['10.0']['C_em_menos8V'], 0)} % na acumulação. Um substrato tipo p, ou um poço p, isolaria a porta "
          "por junção. Vale confirmar o tipo do substrato antes de fixar o leiaute, ou manter óxido espesso fora da porta, "
          "como o artigo já sugere para a inspeção óptica.")

    # ------------------------------------------------------------------ módulo 5
    rel.titulo("Módulo 5, incerteza")
    mcs = R["montecarlo"]
    rel.p("O Monte Carlo usa 1000 amostras por combinação de ramos. As entradas amostradas são energia e dose da "
          "implantação, óxido nativo, limiar de amorfização, solubilidade, aceleração da SPE, temperaturas da SPE e do RTA, "
          "carga fixa, poças, dopagem residual e mobilidade do grafeno e velocidade de Fermi. Os ramos nunca são misturados. "
          "A oxidação usa 900 C com o tempo que dá 20 nm em cada ramo de aceleração, como faria quem escolhe o tempo por um "
          "dos modelos e depois oxida. Nos ramos de fator constante a espessura sai sem dispersão, porque as constantes "
          "cinéticas publicadas não trazem incerteza e a temperatura e o tempo são entradas de varredura. A incerteza real "
          "da espessura está na diferença entre ramos, e só a medida numa amostra de controle a reduz.")
    linhas = []
    for k, m in mcs.items():
        linhas.append([k.replace("|", " "), intervalo(m, "espessura_ox_nm"), intervalo(m, "Rs_final"),
                       intervalo(m, "V_D"), intervalo(m, "perda_si_5e12", lambda v: 100 * v, 2)])
    rel.tabela(["ramos", "óxido (nm)", "Rs da porta (ohm/q)", "V_D (V)", "perda no Si a 5e12 (%)"], linhas,
               "Mediana e intervalo de 95 % do Monte Carlo por combinação de ramos. O sufixo n ou p é o sinal da dopagem residual.")
    from gfetsim import figuras as fg
    mcd = {}
    for k in mcs:
        dd = np.load(RES / f"mc_{k.replace('|', '_')}.npz")
        mcd[k] = {"saidas": {c[2:]: dd[c] for c in dd.files if c.startswith("s_")}}
    fg.intervalos(mcd, [("espessura_ox_nm", "óxido (nm)"), ("Rs_final", "Rs da porta (ohm/q)"), ("V_D", "V_D (V)")],
                  RES / "fig_mc_intervalos.png")
    rel.figura(RES / "fig_mc_intervalos.png", "Mediana (ponto) e intervalo de 95 % (barra) por combinação de ramos. Os ramos "
               "B850 e B900 usam o mesmo tempo de oxidação, calculado para B900, o que mostra o efeito de trocar a variante "
               "de recozimento sem recalcular o tempo.", 16)
    rel.figura(RES / "fig_mc_VD.png", "Distribuição da tensão de Dirac para dopagem residual tipo n e tipo p.", 11)
    sob = R["sobol"]
    legivel = {"espessura_ox_nm": "a espessura do óxido", "Rs_final": "a resistência de folha da porta",
               "V_D": "a tensão de Dirac", "perda_si_5e12": "a perda de capacitância no Si",
               "Rsq_grafeno_5e12": "a resistência de folha do grafeno", "C_min_sobre_Cox": "a capacitância mínima",
               "energia": "energia", "dose": "dose", "oxido_nativo": "óxido nativo", "dose_limiar": "limiar de amorfização",
               "fator_ne": "solubilidade elétrica", "fator_spe": "aceleração da SPE", "T_rta": "temperatura do RTA",
               "T_spe": "temperatura da SPE", "Q_f": "carga fixa", "n_estrela": "poças de carga", "n_dop": "dopagem residual",
               "vF": "velocidade de Fermi", "mu_grafeno": "mobilidade do grafeno"}
    frases = []
    for s, v in sob.items():
        ST = np.array(v["ST"])
        ordem = [i for i in np.argsort(ST)[::-1][:2] if ST[i] > 0.05]
        frases.append(f"{legivel[s][0].upper() + legivel[s][1:]} depende sobretudo de " +
                      " e de ".join(f"{legivel[v['nomes'][i]]} ({br(min(ST[i], 1.0), 2)})" for i in ordem))
    rel.p("Os índices totais de Sobol, calculados no ramo de Ho e Plummer com dopagem tipo p e 960 avaliações da cadeia, "
          "mostram de onde vem a incerteza de cada saída. " + ". ".join(frases) + ". Com esse número de avaliações os "
          "intervalos de confiança dos índices ficam em torno de 0,1 a 0,3, então só a ordem dos fatores dominantes é "
          "confiável, e índices acima de 1 são flutuação de amostragem.")
    rel.figura(RES / "fig_sobol_V_D.png", "Índices de Sobol de primeira ordem e totais para a tensão de Dirac.", 11)

    # ------------------------------------------------------------------ módulo 6
    rel.titulo("Módulo 6, calibração")
    sp = R["sputtering"]
    rel.p(f"O módulo reproduz a regressão do artigo do usuário a partir da Tabela 1. Inclinação de {br(sp['inclinacao'], 3)} "
          f"nm por minuto por watt, com intervalo de confiança de {br(sp['ic95'][0], 3)} a {br(sp['ic95'][1], 3)}, R2 de "
          f"{br(sp['r2'], 3)}, p de {br(sp['p'], 3)} e tempo de {br(sp['tempo_20nm_50W']['tempo'])} min para 20 nm a 50 W, com "
          f"intervalo de previsão de {br(sp['tempo_20nm_50W']['ip95'][0])} a {br(sp['tempo_20nm_50W']['ip95'][1])} min. Todos os "
          "números da Tabela 2 do artigo saem iguais na última casa publicada.")
    rel.p("Os modelos de arquivo para as medidas estão em dados. Espessura de óxido térmico (ajusta o fator de aceleração "
          "por grupo dopado e não dopado), resistência de folha, C-V e curva de transferência (ajusta V_D, mobilidade, "
          "densidade residual e resistência de contato por mínimos quadrados ponderados e, com emcee, a posterior). Os "
          "testes recuperam os parâmetros verdadeiros de curvas sintéticas com ruído de 1 %.")
    rel.figura(RES / "fig_sputtering.png", "Taxa de deposição por sputtering contra potência, com intervalos de confiança "
               "e de previsão de 95 %, refeita a partir da Tabela 1 do artigo.", 9)

    # ------------------------------------------------------------------ limites
    rel.titulo("O que o modelo não captura")
    rel.lista([
        "Canalização. O IMSIL não está disponível e o ângulo de rotação não foi informado. Com 3e15 cm-2 a superfície "
        "amorfiza cedo, o que reduz a canalização, mas a cauda abaixo de 1e18 cm-3 pode ser mais funda que a calculada.",
        "TED. Falta o prefator da dissolução dos defeitos {311}. Afeta a cauda e a profundidade de junção, pouco a superfície.",
        "Aglomerados P-defeito. Sem eles o modelo superestima a ativação em recozimentos de uma etapa.",
        "Cauda acelerada de Fair e Tsai e difusão acelerada pela oxidação.",
        "Estreitamento de banda e efeitos de ionização incompleta no Si n+ (desprezíveis muito acima de Mott, segundo R35).",
        "Efeitos bidimensionais na borda da máscara, na janela de contato da porta e na isolação do substrato.",
        "Ilhas multicamada (17 % da área no artigo), trincas e resíduos de PMMA no grafeno.",
        "Histerese por água. O modelo de relaxação existe (R47), mas seus dois parâmetros só saem de medidas de ida e volta.",
        "Diferença de função trabalho entre grafeno e Si n+, que fica dentro da calibração de V_D.",
    ])

    # ------------------------------------------------------------------ recomendações
    rel.titulo("O que os resultados sugerem para o processo")
    rel.lista([
        "A variante A dá junção mais rasa e porta totalmente ativa, coberta pela validação de R13. A variante B aprofunda "
        "a junção para cerca de 150 a 220 nm e não tem dado direto de ativação.",
        "Para o óxido térmico, 900 a 950 C deixa o tempo entre alguns minutos e uma hora, controlável num forno. A 1000 C "
        "o ramo de Ho e Plummer pede menos de 3 min, curto demais para controlar.",
        "A espessura do óxido térmico sobre o n+ não pode ser tirada de lâmina sem implantação. Uma amostra de controle "
        "com a mesma implantação e o mesmo recozimento resolve a maior incerteza do processo.",
        "O óxido por sputtering evita o consumo de Si, a segregação e a aceleração por dopagem, e sua incerteza vem do "
        "intervalo de previsão do artigo, cerca de 15 a 25 nm para 5,1 min a 50 W.",
        "Medir V_D logo após a transferência define se a operação fica longe da ruptura.",
        "Confirmar o tipo do substrato, porque n+ em substrato n não isola a porta.",
    ])

    # ------------------------------------------------------------------ provisórios
    rel.titulo("Valores provisórios usados")
    rel.p("Os valores abaixo entraram nos cálculos como [I], porque a fonte primária não pôde ser lida. Cada um está no "
          "params.yaml com nota e no pendencias.md com a fonte que resolve.")
    rel.lista(R["provisorios_usados"])

    rel.titulo("Referências citadas")
    rel.p("A tabela completa, com a situação de cada conferência, está em docs/referencias.md, e as notas de conferência "
          "estão em docs/notas_conferencia.md. A lista abaixo traz só as referências citadas neste relatório.")
    citadas = ["R12", "R12b", "R13", "R14", "R19", "R20", "R23", "R25", "R26", "R29", "R32", "R33", "R34", "R35", "R39",
               "R40", "R42", "R45", "R47", "N10", "N12", "N17", "N22", "N24", "A2", "A3", "A4", "A6", "B1", "B2", "G1", "K1",
               "D1", "S1", "S2", "IH", "ART"]
    linhas = {}
    for ln in (RAIZ / "docs" / "referencias.md").read_text().splitlines():
        if ln.startswith("| ") and not ln.startswith("| ID") and not ln.startswith("|---"):
            cel = [c.strip() for c in ln.strip("|").split("|")]
            if cel[0] in citadas and cel[0] not in linhas:
                linhas[cel[0]] = [cel[0], cel[1].replace(":", ","), cel[2]]
    rel.tabela(["ID", "referência", "DOI"], [linhas[k] for k in citadas if k in linhas], "Referências citadas.")

    docx = SAIDA / "relatorio_gfetsim.docx"
    rel.salvar(docx)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(SAIDA), str(docx)], check=True,
                   capture_output=True)
    print(docx, SAIDA / "relatorio_gfetsim.pdf")


if __name__ == "__main__":
    main()
